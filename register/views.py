#
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import user_passes_test, login_required
from django.contrib.auth import login, logout
from django.contrib.auth.forms import AuthenticationForm
from django.contrib import messages
from .models import Intern, Assessment, AssessmentMarks
from .forms import InternForm, AssessmentForm, AssessmentMarksForm
from .staff_form import StaffCreationForm
from django.db.models import Q
from django.http import HttpResponse, JsonResponse
from django.db import transaction
from django.contrib.auth.models import User
import uuid
import openpyxl
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from django.contrib.auth.forms import UserChangeForm


# ------------------- CUSTOM ROLE CHECKERS ------------------- #
def is_superuser(user):
    return user.is_authenticated and user.is_superuser

def is_staff_or_superuser(user):
    # Consolidates both original staff_user definitions
    return user.is_authenticated and (user.is_staff or user.is_superuser)


# ------------------- PUBLIC VIEWS ------------------- #
def intern_form(request):
    if request.method == 'POST':
        form = InternForm(request.POST, request.FILES)
        if form.is_valid():
            intern = form.save()
            return redirect('success', intern_id=intern.id)
    else:
        form = InternForm()
    return render(request, 'register/form.html', {'form': form})

def success_view(request, intern_id):
    intern = get_object_or_404(Intern, id=intern_id)
    return render(request, 'register/success.html', {'intern': intern})


# ------------------- ADMIN LOGIN & LOGOUT ------------------- #
def admin_login(request):
    if is_staff_or_superuser(request.user):
        return redirect('dashboard_home')

    if request.method == 'POST':
        form = AuthenticationForm(data=request.POST)
        if form.is_valid():
            user = form.get_user()
            if is_staff_or_superuser(user):
                login(request, user)
                return redirect('dashboard_home')
            else:
                messages.error(request, "Access denied: not a staff or superuser.")
        else:
            messages.error(request, "Invalid username or password.")
    else:
        form = AuthenticationForm()
    return render(request, 'register/admin_login.html', {'form': form})

@user_passes_test(is_staff_or_superuser)
def admin_logout(request):
    logout(request)
    return redirect('admin_login')


# ------------------- STAFF + SUPERUSER: DASHBOARD ------------------- #
@login_required
@user_passes_test(is_staff_or_superuser)
def admin_dashboard(request):
    query = request.GET.get('q')
    interns = Intern.objects.all().order_by('id')
    if query:
        interns = interns.filter(
            Q(full_name__icontains=query) |
            Q(email__icontains=query) |
            Q(qualification__icontains=query)
        )
    if request.method == 'POST':
        form = InternForm(request.POST, request.FILES)
        if form.is_valid():
            form.save()
            messages.success(request, "Intern application submitted successfully.")
            return redirect('admin_dashboard')
    else:
        form = InternForm()

    return render(request, 'register/admin_dashboard.html', {
        'interns': interns,
        'form': form,
        'query': query,
        'is_superuser': request.user.is_superuser,
        'is_mentor': request.user.groups.filter(name='mentor').exists(),
        'is_employee': request.user.groups.filter(name='employee').exists(),
    })


# ------------------- SUPERUSER-ONLY VIEWS ------------------- #
@user_passes_test(is_superuser)
def admin_edit_intern(request, intern_id):
    intern = get_object_or_404(Intern, id=intern_id)
    if request.method == 'POST':
        form = InternForm(request.POST, request.FILES, instance=intern)
        if form.is_valid():
            form.save()
            messages.success(request, "Intern updated successfully.")
            return redirect('admin_dashboard')
    else:
        form = InternForm(instance=intern)
    return render(request, 'register/admin_edit_form.html', {'form': form, 'intern': intern})

@user_passes_test(is_superuser)
def admin_delete_intern(request, intern_id):
    intern = get_object_or_404(Intern, id=intern_id)
    if request.method == 'POST':
        intern.delete()
        messages.success(request, "Intern deleted successfully.")
        return redirect('admin_dashboard')
    return render(request, 'register/confirm_delete.html', {'intern': intern})

@user_passes_test(is_superuser)
def add_staff_user(request):
    if request.method == 'POST':
        form = StaffCreationForm(request.POST)
        if form.is_valid():
            form.save()
            messages.success(request, "New staff user created successfully!")
            return redirect('admin_dashboard')
    else:
        form = StaffCreationForm()
    return render(request, 'register/add_staff.html', {'form': form})

@user_passes_test(is_superuser)
def create_assessment(request):
    if request.method == "POST":
        assessment_form = AssessmentForm(request.POST)
        marks_form = AssessmentMarksForm(request.POST)
        if assessment_form.is_valid() and marks_form.is_valid():
            try:
                with transaction.atomic():
                    assessment = assessment_form.save()
                    marks = marks_form.save(commit=False)
                    marks.assessment = assessment
                    marks.save()
                return redirect("dashboard")
            except Exception as e:
                messages.error(request, f"Error saving assessment: {e}")
        else:
            messages.error(request, "Please correct the errors below.")
    else:
        assessment_form = AssessmentForm()
        marks_form = AssessmentMarksForm()
    context = {
        "assessment_form": assessment_form,
        "marks_form": marks_form,
    }
    return render(request, "register/assessment_form.html", context)

@user_passes_test(is_superuser)
def trainee_domain_assessments_json(request, trainee_id, domain):
    trainee = get_object_or_404(Intern, id=trainee_id)
    qs = Assessment.objects.filter(trainee_name=trainee, domain=domain).order_by("-assessment_date", "-id")
    data = [{
        "id": a.id,
        "assessor_name": a.assessor_name,
        "assessment_date": a.assessment_date.strftime("%Y-%m-%d") if a.assessment_date else "",
        "topics": a.topics,
        "assessment_number": a.assessment_number,
    } for a in qs]
    return JsonResponse({
        "trainee": trainee.full_name,
        "domain": domain,
        "count": len(data),
        "assessments": data
    })

@user_passes_test(is_superuser)
def edit_assessment_step1(request, assessment_id: uuid.UUID):
    assessment = get_object_or_404(Assessment, id=assessment_id)
    if request.method == "POST":
        form = AssessmentForm(request.POST, instance=assessment)
        if form.is_valid():
            form.save()
            return redirect("edit_assessment_step2", assessment_id=assessment.id)
        else:
            print("Form is not valid. Errors:", form.errors)
            return render(request, "register/edit_assessment_step1.html", {"form": form, "assessment": assessment})
    else:
        form = AssessmentForm(instance=assessment)
    return render(request, "register/edit_assessment_step1.html", {"form": form, "assessment": assessment})

@user_passes_test(is_superuser)
def edit_assessment_step2(request, assessment_id: uuid.UUID):
    assessment = get_object_or_404(Assessment, id=assessment_id)
    marks, created = AssessmentMarks.objects.get_or_create(assessment=assessment)
    if request.method == "POST":
        form = AssessmentMarksForm(request.POST, instance=marks)
        if form.is_valid():
            form.save()
            return redirect("dashboard")
    else:
        form = AssessmentMarksForm(instance=marks)
    return render(request, "register/edit_assessment_step2.html", {"form": form, "assessment": assessment})

@user_passes_test(is_superuser)
def delete_assessment(request, trainee_id):
    assessments_to_delete = Assessment.objects.filter(trainee_name_id=trainee_id)
    if request.method == 'POST':
        assessments_to_delete.delete()
        messages.success(request, 'All assessments for the trainee were successfully deleted!')
        return redirect('dashboard')
    trainee = get_object_or_404(Intern, id=trainee_id)
    return render(request, 'register/confirm_delete.html', {'trainee': trainee})

@user_passes_test(is_superuser)
def edit_staff_user(request, user_id):
    user = get_object_or_404(User, id=user_id)
    if request.method == 'POST':
        form = UserChangeForm(request.POST, instance=user)
        if form.is_valid():
            form.save()
            return redirect('admin_user')
    else:
        form = UserChangeForm(instance=user)
    return render(request, 'register/edit_staff_user.html', {'form': form, 'user': user})

# ------------------- DATA EXPORT HELPERS ------------------- #
def export_to_excel(filename, headers, data_rows):
    """A reusable helper function to generate and return an Excel file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = filename.split('.')[0]

    ws.append(headers)
    for row in data_rows:
        ws.append(row)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response

# ------------------- SUPERADMIN, ADMIN LISTS & PROFILES ------------------- #
def super_admin_view(request):
    super_admins = User.objects.filter(is_superuser=True)
    return render(request, "register/super_admin.html", {"super_admins": super_admins})

def admin_user_view(request):
    admins = User.objects.filter(is_staff=True, is_superuser=False)
    return render(request, "register/admin_user_partial.html", {"admins": admins})

@login_required
def profile_view(request):
    return render(request, "register/profile.html", {"user": request.user})

@login_required
def dashboard_home(request):
    return render(request, "register/dashboard_home.html")


# ------------------- ASSESSMENT DASHBOARD & RELATED VIEWS ------------------- #
def dashboard(request):
    query = request.GET.get('q')
    assessments_queryset = Assessment.objects.all()

    if query:
        assessments_queryset = assessments_queryset.filter(
            Q(trainee_name__full_name__icontains=query) |
            Q(domain__icontains=query) |
            Q(assessor_name__icontains=query)
        )

    groups = assessments_queryset.values("trainee_name__id", "trainee_name__full_name", "domain").distinct()
    data = []

    for g in groups:
        trainee_id = g["trainee_name__id"]
        trainee_name = g["trainee_name__full_name"]
        domain = g["domain"]

        assessments = Assessment.objects.filter(trainee_name_id=trainee_id, domain=domain).order_by("assessment_number")[:3]

        marks_list = [a.total_marks for a in assessments]
        assessors = {a.assessor_name for a in assessments if a.assessor_name}

        total_scored = sum(m for m in marks_list if m is not None)
        no_of_assessments = len([m for m in marks_list if m is not None])

        while len(marks_list) < 3:
            marks_list.append(None)

        max_possible = no_of_assessments * Assessment.MAX_MARKS
        percentage = round((total_scored / max_possible) * 100, 2) if max_possible > 0 else 0

        data.append({
            "trainee_id": trainee_id,
            "trainee": trainee_name,
            "domain": domain,
            "assessor": ", ".join(assessors) if assessors else "-",
            "assessment1": marks_list[0],
            "assessment2": marks_list[1],
            "assessment3": marks_list[2],
            "total_marks": total_scored,
            "percentage": percentage,
        })
    return render(request, "register/dashboard.html", {"data": data})

def trainee_assessments(request, trainee_id):
    trainee = get_object_or_404(Intern, id=trainee_id)
    assessments = Assessment.objects.filter(trainee_name=trainee).order_by('created_at')
    
    assessments_with_marks = []
    for a in assessments:
        marks = AssessmentMarks.objects.filter(assessment=a).first()
        assessments_with_marks.append({
            'assessment': a,
            'marks': marks
        })
    
    return render(request, 'register/traine_assessments.html', {
        'trainee': trainee,
        'assessments_with_marks': assessments_with_marks
    })
    


# ------------------- EXPORT VIEWS ------------------- #
@user_passes_test(is_staff_or_superuser)
def export_dashboard_excel(request):
    headers = [
        "Trainee Name", "Domain", "Assessor",
        "Assessment 1", "Assessment 2", "Assessment 3",
        "Total Marks", "Percentage"
    ]
    data_rows = []
    
    # Get all interns who have at least one assessment
    interns = Intern.objects.filter(assessments__isnull=False).distinct()
    
    for intern in interns:
        assessments = Assessment.objects.filter(trainee_name=intern)
        
        # Use a dictionary to store marks and handle missing assessments gracefully
        assessment_marks = {a.assessment_number: a.total_marks for a in assessments}
        a1 = assessment_marks.get(1, 0)
        a2 = assessment_marks.get(2, 0)
        a3 = assessment_marks.get(3, 0)

        total = sum(m for m in [a1, a2, a3] if m is not None)
        
        # Calculate percentage based on number of completed assessments
        no_of_assessments = len(assessment_marks)
        max_total = no_of_assessments * Assessment.MAX_MARKS if no_of_assessments > 0 else 0
        percentage = (total / max_total) * 100 if max_total > 0 else 0
        
        row = [
            intern.full_name,
            assessments.first().domain if assessments.exists() else "",
            assessments.first().assessor_name if assessments.exists() else "",
            a1, a2, a3,
            total,
            round(percentage, 2),
        ]
        data_rows.append(row)
        
    return export_to_excel("dashboard.xlsx", headers, data_rows)


@user_passes_test(is_staff_or_superuser)
def export_trainee_excel(request, trainee_id):
    trainee = get_object_or_404(Intern, id=trainee_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Trainee Report"

    # --- Trainee Info ---
    ws.append(["Trainee Information"])
    ws.append(["Name", trainee.full_name])
    ws.append(["Domain / Area", trainee.area_of_interest])

    first_assessment = Assessment.objects.filter(trainee_name=trainee,domain=trainee.area_of_interest ).order_by("assessment_number").first()
    assessor_name = first_assessment.assessor_name if first_assessment else "Not Assigned"
    ws.append(["Assessor Name", assessor_name])
    ws.append([])

    # --- Overall Performance ---
    assessments = Assessment.objects.filter(trainee_name=trainee, domain=trainee.area_of_interest   ).order_by("assessment_number")

    overall_total = sum(a.total_marks or 0 for a in assessments)
    total_assessments_count = assessments.count()
    max_total_possible = total_assessments_count * Assessment.MAX_MARKS
    overall_percentage = (overall_total / max_total_possible * 100) if max_total_possible > 0 else 0

    ws.append(["Overall Performance"])
    ws.append(["Total Marks Obtained", overall_total])
    ws.append(["Overall Percentage", f"{overall_percentage:.2f}%"])
    ws.append([])

    # --- Assessment Marks ---
    ws.append(["Assessment Marks"])
    headers = [
        "Assessment #", "Technical", "Analytical", "Troubleshooting",
        "Verbal", "Written", "Collaboration", "Time Mgmt", "Work Ethics",
        "Total Marks", "Percentage", "Date"
    ]
    ws.append(headers)

    for i, assessment in enumerate(assessments, start=1):
        # take first marks entry (or aggregate multiple if needed)
        marks = assessment.marks.first()

        if marks:
            scores = [
                marks.technical_skill, marks.analytical_skills, marks.troubleshooting,
                marks.verbal_communication, marks.written_communication,
                marks.collaboration, marks.time_management, marks.work_ethic
            ]
            total = marks.total_marks
            percentage = round((total / Assessment.MAX_MARKS) * 100, 2)
        else:
            scores = [0] * 8
            total = 0
            percentage = 0.0

        ws.append([
            f"Assessment {i}", *scores, total, f"{percentage}%",
            assessment.assessment_date.strftime("%Y-%m-%d") if assessment.assessment_date else ""
        ])

    # --- Auto column width ---
    for col in ws.columns:
        max_length = 0
        column = col[0].column
        for cell in col:
            try:
                if cell.value and len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        adjusted_width = max_length + 2
        ws.column_dimensions[get_column_letter(column)].width = adjusted_width

    # --- Response ---
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{trainee.full_name}_Assessment_Report.xlsx"'
    wb.save(response)
    return response





